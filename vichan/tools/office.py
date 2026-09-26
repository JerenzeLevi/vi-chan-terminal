"""Office tools: Excel (.xlsx), Word (.docx), PowerPoint (.pptx)."""
from pathlib import Path

from . import tool


def _out(path):
    p = Path(path)
    p.parent.mkdir(parents=True, exist_ok=True)
    return p


# ---------------------------------------------------------------- Excel
@tool(
    "create_xlsx",
    "Create an Excel workbook. Each sheet has a name and rows (list of lists). Strings starting with '=' become formulas. "
    "The first row is styled as a bold header.",
    {
        "path": {"type": "string"},
        "sheets": {
            "type": "array",
            "items": {
                "type": "object",
                "properties": {
                    "name": {"type": "string"},
                    "rows": {"type": "array", "items": {"type": "array", "items": {}}},
                },
                "required": ["name", "rows"],
            },
        },
    },
    ["path", "sheets"],
    approval=True,
)
def create_xlsx(path, sheets):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    if not sheets:
        return "error: give at least one sheet"
    wb = Workbook()
    wb.remove(wb.active)
    for i, sh in enumerate(sheets):
        ws = wb.create_sheet(str(sh.get("name") or f"Sheet{i + 1}")[:31])
        rows = sh.get("rows", [])
        for r in rows:
            ws.append(r)
        if rows:
            for c in ws[1]:
                c.font = Font(bold=True)
        for j, col in enumerate(ws.columns, 1):
            width = max((len(str(c.value)) for c in col if c.value is not None and not str(c.value).startswith("=")), default=8)
            ws.column_dimensions[get_column_letter(j)].width = min(width + 2, 60)
    wb.save(_out(path))
    return f"created {path} with {len(sheets)} sheet(s)"


@tool(
    "read_xlsx",
    "Read an Excel workbook. Formulas are shown as formula text. Returns tab-separated rows per sheet.",
    {"path": {"type": "string"}, "sheet": {"type": "string"}, "max_rows": {"type": "integer"}},
    ["path"],
    readonly=True,
)
def read_xlsx(path, sheet="", max_rows=200):
    from openpyxl import load_workbook

    wb = load_workbook(path)
    out = []
    for ws in wb.worksheets:
        if sheet and ws.title != sheet:
            continue
        out.append(f"## {ws.title} ({ws.max_row} rows x {ws.max_column} cols)")
        for row in ws.iter_rows(max_row=max_rows, values_only=True):
            out.append("\t".join("" if v is None else str(v) for v in row))
    return "\n".join(out)[:30000] or "no matching sheet"


@tool(
    "update_xlsx_cells",
    "Set cell values in an existing workbook. cells maps addresses like 'B2' to values ('=SUM(A1:A9)' for formulas).",
    {"path": {"type": "string"}, "cells": {"type": "object"}, "sheet": {"type": "string"}},
    ["path", "cells"],
    approval=True,
)
def update_xlsx_cells(path, cells, sheet=""):
    from openpyxl import load_workbook

    wb = load_workbook(path)
    ws = wb[sheet] if sheet else wb.active
    for addr, val in cells.items():
        ws[addr] = val
    wb.save(path)
    return f"updated {len(cells)} cell(s) in {ws.title}"


# ---------------------------------------------------------------- Word
@tool(
    "create_docx",
    "Create a Word document from blocks. Block types: heading{text,level}, paragraph{text,bold,italic}, "
    "bullets{items}, numbered{items}, table{rows}, image{path}, page_break.",
    {
        "path": {"type": "string"},
        "blocks": {
            "type": "array",
            "items": {
                "type": "object",
                "properties": {
                    "type": {"type": "string", "enum": ["heading", "paragraph", "bullets", "numbered", "table", "image", "page_break"]},
                    "text": {"type": "string"},
                    "level": {"type": "integer"},
                    "bold": {"type": "boolean"},
                    "italic": {"type": "boolean"},
                    "items": {"type": "array", "items": {"type": "string"}},
                    "rows": {"type": "array", "items": {"type": "array", "items": {}}},
                    "path": {"type": "string"},
                },
                "required": ["type"],
            },
        },
    },
    ["path", "blocks"],
    approval=True,
)
def create_docx(path, blocks):
    from docx import Document
    from docx.shared import Inches

    doc = Document()
    for b in blocks:
        t = b.get("type")
        if t == "heading":
            doc.add_heading(b.get("text", ""), level=max(0, min(int(b.get("level", 1)), 9)))
        elif t == "paragraph":
            run = doc.add_paragraph().add_run(b.get("text", ""))
            run.bold, run.italic = bool(b.get("bold")), bool(b.get("italic"))
        elif t in ("bullets", "numbered"):
            style = "List Bullet" if t == "bullets" else "List Number"
            for item in b.get("items", []):
                doc.add_paragraph(str(item), style=style)
        elif t == "table":
            rows = b.get("rows", [])
            if rows:
                table = doc.add_table(rows=len(rows), cols=max(len(r) for r in rows))
                table.style = "Table Grid"
                for i, r in enumerate(rows):
                    for j, v in enumerate(r):
                        cell = table.cell(i, j)
                        cell.text = "" if v is None else str(v)
                        if i == 0:
                            for run in cell.paragraphs[0].runs:
                                run.bold = True
        elif t == "image":
            doc.add_picture(b["path"], width=Inches(6))
        elif t == "page_break":
            doc.add_page_break()
    doc.save(_out(path))
    return f"created {path} ({len(blocks)} blocks)"


@tool("read_docx", "Read a Word document: paragraphs (with style) and tables.", {"path": {"type": "string"}}, ["path"], readonly=True)
def read_docx(path):
    from docx import Document

    doc = Document(path)
    out = [f"[{p.style.name}] {p.text}" if p.style.name != "Normal" else p.text for p in doc.paragraphs if p.text.strip()]
    for n, table in enumerate(doc.tables, 1):
        out.append(f"## table {n}")
        out += ["\t".join(c.text for c in row.cells) for row in table.rows]
    return "\n".join(out)[:30000] or "(empty document)"


# ---------------------------------------------------------------- PowerPoint
@tool(
    "create_pptx",
    "Create a 16:9 PowerPoint deck. Each slide: title, optional subtitle (title slide), bullets (list of strings; "
    "prefix with two spaces per indent level), notes (speaker notes), image (file path).",
    {
        "path": {"type": "string"},
        "slides": {
            "type": "array",
            "items": {
                "type": "object",
                "properties": {
                    "title": {"type": "string"},
                    "subtitle": {"type": "string"},
                    "bullets": {"type": "array", "items": {"type": "string"}},
                    "notes": {"type": "string"},
                    "image": {"type": "string"},
                },
                "required": ["title"],
            },
        },
    },
    ["path", "slides"],
    approval=True,
)
def create_pptx(path, slides):
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    full = prs.slide_width - Inches(1.4)
    for spec in slides:
        bullets, image = spec.get("bullets") or [], spec.get("image")
        layout = 0 if (spec.get("subtitle") and not bullets) else 1 if bullets else 5
        s = prs.slides.add_slide(prs.slide_layouts[layout])
        for ph in s.placeholders:  # the default template is 4:3, so re-fit placeholders to 16:9
            ph.left, ph.width = Inches(0.7), full
        s.shapes.title.text = spec.get("title", "")
        if layout == 0:
            s.placeholders[1].text = spec["subtitle"]
        if bullets:
            body = s.placeholders[1]
            if image:
                body.width = Inches(6.2)
            tf = body.text_frame
            for i, b in enumerate(bullets):
                para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                para.text = b.lstrip()
                para.level = min((len(b) - len(b.lstrip())) // 2, 4)
        if image and Path(image).is_file():
            if bullets:
                s.shapes.add_picture(image, Inches(7.3), Inches(1.8), width=Inches(5.4))
            else:
                s.shapes.add_picture(image, Inches(2.2), Inches(1.7), height=Inches(5.3))
        if spec.get("notes"):
            s.notes_slide.notes_text_frame.text = spec["notes"]
    prs.save(_out(path))
    return f"created {path} with {len(slides)} slide(s)"


@tool("read_pptx", "Read a PowerPoint deck: text on every slide plus speaker notes.", {"path": {"type": "string"}}, ["path"], readonly=True)
def read_pptx(path):
    from pptx import Presentation

    out = []
    for i, s in enumerate(Presentation(path).slides, 1):
        out.append(f"## slide {i}")
        out += [sh.text_frame.text for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
        if s.has_notes_slide and s.notes_slide.notes_text_frame.text.strip():
            out.append(f"(notes) {s.notes_slide.notes_text_frame.text}")
    return "\n".join(out)[:30000] or "(empty deck)"


# ---------------------------------------------------------------- edit text in Word / PowerPoint
def _replace_in_paragraph(p, old, new):
    """Replace text in one paragraph; keeps run formatting when the text sits inside a single run."""
    if old not in p.text:
        return 0
    n = p.text.count(old)
    for r in p.runs:
        if old in r.text:
            r.text = r.text.replace(old, new)
    if old in p.text:  # the text spans several runs: merge into the first run
        full = p.text.replace(old, new)
        for i, r in enumerate(p.runs):
            r.text = full if i == 0 else ""
    return n


def _shape_paragraphs(shape):
    if shape.has_text_frame:
        yield from shape.text_frame.paragraphs
    if getattr(shape, "has_table", False) and shape.has_table:
        for row in shape.table.rows:
            for cell in row.cells:
                yield from cell.text_frame.paragraphs
    if shape.shape_type == 6:  # group
        for sub in shape.shapes:
            yield from _shape_paragraphs(sub)


@tool(
    "edit_office_text",
    "Find-and-replace text inside an existing Word (.docx) or PowerPoint (.pptx) file, including tables. "
    "Replaces every occurrence. Use read_docx / read_pptx first to see the exact text.",
    {"path": {"type": "string"}, "old_text": {"type": "string"}, "new_text": {"type": "string"}},
    ["path", "old_text", "new_text"],
    approval=True,
)
def edit_office_text(path, old_text, new_text):
    ext = Path(path).suffix.lower()
    count = 0
    if ext == ".docx":
        from docx import Document

        doc = Document(path)
        paras = list(doc.paragraphs)
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    paras += cell.paragraphs
        for sec in doc.sections:
            paras += sec.header.paragraphs + sec.footer.paragraphs
        for p in paras:
            count += _replace_in_paragraph(p, old_text, new_text)
        doc.save(path)
    elif ext == ".pptx":
        from pptx import Presentation

        prs = Presentation(path)
        for slide in prs.slides:
            for shape in slide.shapes:
                for p in _shape_paragraphs(shape):
                    count += _replace_in_paragraph(p, old_text, new_text)
            if slide.has_notes_slide:
                for p in slide.notes_slide.notes_text_frame.paragraphs:
                    count += _replace_in_paragraph(p, old_text, new_text)
        prs.save(path)
    else:
        return "error: only .docx and .pptx are supported (use edit_file for text files, update_xlsx_cells for Excel)"
    return f"replaced {count} occurrence(s) in {path}" if count else "error: text not found"
